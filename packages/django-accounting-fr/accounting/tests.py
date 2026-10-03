"""Tests de l'app comptable, indépendants de tout projet hôte (seule l'API publique est utilisée)."""
from datetime import date
from decimal import Decimal
import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import api, exports, reports
from .models import Account, AccountingSettings, FiscalPeriod, Journal, LedgerEntry, PaymentAccount, Transaction, VATRate
from .posting import (
    AccountingError, Line, close_period, generate_opening_entries, post_entry, reverse_entry, validate_entry,
)

TODAY = timezone.localdate()
CUSTOMER = api.Party("C001", "Dupont SARL")


def acc(code):
    return Account.objects.get(code=code)


def balanced_lines(amount="100.00"):
    return [Line(acc("411000"), debit=Decimal(amount)), Line(acc("707000"), credit=Decimal(amount))]


def accountant(email="compta@example.com", superuser=False):
    User = get_user_model()
    user = User.objects.create_user(username=email, email=email, password="x" * 12, is_superuser=superuser)
    user.user_permissions.add(*Permission.objects.filter(content_type__app_label="accounting"))
    return user


def sale(key="invoice:1", amount="120.00", rate="0.2000", day=TODAY):
    return api.post_sale(key, day, "F-1", "Facture F-1", CUSTOMER,
                         [api.SaleLine(Decimal(amount), Decimal(rate))], document=("invoice", key.split(":")[-1]))


def balances(txn):
    result = {}
    for e in txn.entries.all():
        result[e.account.code] = result.get(e.account.code, Decimal("0")) + e.debit - e.credit
    return result


class ChartTest(TestCase):
    def test_default_chart_is_installed(self):
        self.assertTrue(Account.objects.filter(code="411000").exists())
        self.assertEqual(set(Journal.objects.values_list("code", flat=True)), {"VT", "BQ", "ST", "PP", "CA", "AC", "OD", "AN"})
        self.assertEqual(Journal.objects.get(code="BQ").account.code, "512000")
        self.assertEqual(VATRate.objects.get(rate=Decimal("0.2000")).collected_account.code, "445711")
        self.assertEqual(AccountingSettings.get().payment_account("stripe").code, "467100")
        self.assertEqual(AccountingSettings.get().payment_account("inconnu").code, "512000")  # repli : banque

    def test_init_accounts_is_idempotent(self):
        from django.core.management import call_command
        before = Account.objects.count()
        call_command("init_accounts", stdout=io.StringIO())
        self.assertEqual(Account.objects.count(), before)


class PostingTest(TestCase):
    def test_unbalanced_entry_is_refused(self):
        with self.assertRaises(AccountingError):
            post_entry("OD", TODAY, "Test", [Line(acc("411000"), debit=Decimal("10")), Line(acc("707000"), credit=Decimal("9"))])
        self.assertFalse(Transaction.objects.exists())

    def test_line_must_have_one_side_in_database(self):
        txn = post_entry("OD", TODAY, "Test", balanced_lines())
        with self.assertRaises(IntegrityError), transaction.atomic():
            LedgerEntry.objects.bulk_create([LedgerEntry(transaction=txn, account=acc("411000"), debit=1, credit=1)])

    def test_validation_numbers_sequentially_per_journal(self):
        first = validate_entry(post_entry("OD", TODAY, "A", balanced_lines()))
        second = validate_entry(post_entry("OD", TODAY, "B", balanced_lines()))
        other = validate_entry(post_entry("VT", TODAY, "C", balanced_lines()))
        year = FiscalPeriod.objects.get().name
        self.assertEqual([first.number, second.number, other.number], [f"OD{year}-00001", f"OD{year}-00002", f"VT{year}-00001"])

    def test_validated_entry_is_locked(self):
        txn = post_entry("OD", TODAY, "Test", balanced_lines(), validate=True)
        txn.description = "Modifiée"
        with self.assertRaises(ValidationError):
            txn.save()
        with self.assertRaises(ValidationError):
            txn.delete()
        line = txn.entries.first()
        line.debit = Decimal("1")
        with self.assertRaises(ValidationError):
            line.save()
        with self.assertRaises(ValidationError):
            line.delete()

    def test_reversal_mirrors_lines_once(self):
        txn = post_entry("OD", TODAY, "Test", balanced_lines("42.00"), validate=True)
        reversal = reverse_entry(txn)
        self.assertTrue(reversal.is_validated)
        self.assertEqual(reversal.reversal_of, txn)
        self.assertEqual(acc("411000").get_balance(), Decimal("0.00"))
        self.assertEqual(reverse_entry(txn).pk, reversal.pk)

    def test_closed_period_refuses_entries(self):
        post_entry("OD", TODAY, "A", balanced_lines(), validate=True)
        close_period(FiscalPeriod.objects.get())
        with self.assertRaises(AccountingError):
            post_entry("OD", TODAY, "B", balanced_lines())

    def test_period_with_drafts_cannot_be_closed(self):
        post_entry("OD", TODAY, "A", balanced_lines())
        with self.assertRaises(AccountingError):
            close_period(FiscalPeriod.objects.get())


class ApiTest(TestCase):
    def test_sale_with_mixed_rates_and_kinds(self):
        txn = api.post_sale("invoice:7", TODAY, "F-7", "Facture F-7", CUSTOMER, [
            api.SaleLine(Decimal("60.00"), Decimal("0.2000")),
            api.SaleLine(Decimal("40.00"), Decimal("0.0550")),
            api.SaleLine(Decimal("6.00"), Decimal("0.2000"), kind="shipping"),
        ], document=("invoice", "7"))
        self.assertTrue(txn.is_validated)
        self.assertEqual(balances(txn), {
            "411000": Decimal("106.00"), "707000": Decimal("-87.91"), "708500": Decimal("-5.00"),
            "445711": Decimal("-11.00"), "445713": Decimal("-2.09")})
        customer_line = txn.entries.get(account__code="411000")
        self.assertEqual((customer_line.auxiliary_code, customer_line.auxiliary_label), ("C001", "Dupont SARL"))

    def test_prices_excluding_tax(self):
        txn = api.post_sale("invoice:8", TODAY, "F-8", "F-8", CUSTOMER, [api.SaleLine(Decimal("100.00"))],
                            prices_include_tax=False)
        self.assertEqual(balances(txn)["411000"], Decimal("120.00"))

    def test_payment_uses_method_account(self):
        sale()
        PaymentAccount.objects.create(method="transfer", label="Virement", journal=Journal.objects.get(code="BQ"))
        payment = api.post_payment("invoice:1:payment", TODAY, "VIR-1", "Règlement", Decimal("120.00"), "transfer", CUSTOMER)
        self.assertEqual(balances(payment), {"512000": Decimal("120.00"), "411000": Decimal("-120.00")})
        self.assertEqual(acc("411000").get_balance(), Decimal("0.00"))

    def test_credit_note_and_refund_cancel_everything(self):
        sale()
        api.post_payment("invoice:1:payment", TODAY, "P", "Règlement", Decimal("120.00"), "stripe", CUSTOMER)
        api.post_credit_note("invoice:1:refund", "invoice:1", TODAY, "AV-1", "Avoir")
        api.post_refund("invoice:1:refund_payment", TODAY, "AV-1", "Remboursement", Decimal("120.00"), "stripe", CUSTOMER)
        for code in ("411000", "707000", "445711", "467100"):
            self.assertEqual(acc(code).get_balance(), Decimal("0.00"), code)
        self.assertEqual(reports.profit_and_loss(TODAY.replace(month=1, day=1), TODAY)["total_revenues"], Decimal("0.00"))

    def test_fees_and_payout_reconcile_the_provider_account(self):
        sale()
        api.post_payment("invoice:1:payment", TODAY, "P", "Règlement", Decimal("120.00"), "stripe", CUSTOMER)
        fee = api.post_fee("invoice:1:fee", TODAY, "ch_1", "Frais Stripe", Decimal("1.93"), "stripe")
        self.assertEqual((fee.type, balances(fee)), ("fee", {"627000": Decimal("1.93"), "467100": Decimal("-1.93")}))
        payout = api.post_transfer("payout:po_1", TODAY, "po_1", "Virement Stripe", Decimal("118.07"), "stripe")
        self.assertEqual((payout.journal.code, balances(payout)), ("ST", {"580000": Decimal("118.07"), "467100": Decimal("-118.07")}))
        received = Transaction.objects.get(source_key="payout:po_1:banque")
        self.assertEqual((received.journal.code, balances(received)), ("BQ", {"512000": Decimal("118.07"), "580000": Decimal("-118.07")}))
        for code, balance in (("467100", "0.00"), ("580000", "0.00"), ("512000", "118.07")):
            self.assertEqual(acc(code).get_balance(), Decimal(balance), code)
        self.assertEqual(api.post_transfer("payout:po_1", TODAY, "po_1", "Virement Stripe", Decimal("118.07"), "stripe").pk, payout.pk)
        self.assertEqual(Transaction.objects.filter(type="transfer").count(), 2)
        self.assertEqual(api.post_fee("invoice:1:fee", TODAY, "ch_1", "Frais Stripe", Decimal("1.93"), "stripe").pk, fee.pk)

    def test_negative_transfer_debits_the_provider(self):
        txn = api.post_transfer("payout:po_2", TODAY, "po_2", "Prélèvement Stripe", Decimal("-10.00"), "stripe")
        self.assertEqual(balances(txn), {"467100": Decimal("10.00"), "580000": Decimal("-10.00")})
        self.assertEqual(balances(Transaction.objects.get(source_key="payout:po_2:banque")),
                         {"580000": Decimal("10.00"), "512000": Decimal("-10.00")})

    def test_transfer_from_the_bank_itself_is_refused(self):
        with self.assertRaises(AccountingError):
            api.post_transfer("payout:x", TODAY, "x", "Virement", Decimal("10.00"), "virement-inconnu")

    def test_credit_note_requires_the_sale(self):
        with self.assertRaises(AccountingError):
            api.post_credit_note("x:refund", "x:sale", TODAY, "AV", "Avoir")

    def test_keys_are_idempotent(self):
        first, second = sale(), sale()
        self.assertEqual(first.pk, second.pk)
        self.assertTrue(api.is_posted("invoice:1"))
        self.assertEqual([t.pk for t in api.entries_for("invoice", "1")], [first.pk])

    def test_missing_vat_rate_setup(self):
        VATRate.objects.filter(rate=Decimal("0.1000")).update(collected_account=None)
        with self.assertRaises(AccountingError):
            sale(rate="0.1000")


class ReportsTest(TestCase):
    def setUp(self):
        sale()
        api.post_payment("invoice:1:payment", TODAY, "P", "Règlement", Decimal("120.00"), "stripe", CUSTOMER)
        self.year_start = TODAY.replace(month=1, day=1)

    def test_trial_balance(self):
        balance = reports.trial_balance(self.year_start, TODAY)
        self.assertTrue(balance["is_balanced"])
        self.assertEqual([g["code"] for g in balance["classes"]], ["4", "7"])

    def test_profit_and_loss_includes_today(self):
        self.assertEqual(reports.profit_and_loss(self.year_start, TODAY)["total_revenues"], Decimal("100.00"))

    def test_general_ledger_running_balance(self):
        ledger = reports.general_ledger(acc("411000"), self.year_start, TODAY)
        self.assertEqual([r["balance"] for r in ledger["rows"]], [Decimal("120.00"), Decimal("0.00")])

    def test_vat_summary(self):
        vat = reports.vat_summary(self.year_start, TODAY)
        self.assertEqual((vat["rows"][0]["base"], vat["net"]), (Decimal("100.00"), Decimal("20.00")))

    def test_vat_summary_includes_manual_entries_without_rate(self):
        post_entry("VT", TODAY, "Vente saisie à la main", [
            Line(acc("411000"), debit=Decimal("60.00")), Line(acc("707000"), credit=Decimal("50.00")),
            Line(acc("445711"), credit=Decimal("10.00"))], validate=True)
        vat = reports.vat_summary(self.year_start, TODAY)
        self.assertEqual([(r["rate"], r["base"], r["vat"]) for r in vat["rows"]],
                         [(Decimal("0.2000"), Decimal("150.00"), Decimal("30.00"))])
        self.assertEqual(reports.dashboard()["vat_month"], Decimal("30.00"))


class ExportsTest(TestCase):
    def setUp(self):
        sale()
        api.post_payment("invoice:1:payment", TODAY, "P", "Règlement", Decimal("120.00"), "stripe", CUSTOMER)
        post_entry("OD", TODAY, "=HYPERLINK()", balanced_lines("5.00"))  # brouillon
        self.start, self.end = TODAY.replace(month=1, day=1), TODAY.replace(month=12, day=31)

    def test_fec_format(self):
        filename, content = exports.fec_file(self.start, self.end, "123 456 789 00011")
        self.assertEqual(filename, f"123456789FEC{self.end:%Y%m%d}.txt")
        rows = [line.split("\t") for line in content.decode("iso-8859-15").rstrip("\r\n").split("\r\n")]
        self.assertEqual(rows[0], exports.FEC_COLUMNS)
        self.assertTrue(all(len(row) == 18 for row in rows))
        self.assertEqual(len(rows) - 1, 5)  # 3 lignes de vente + 2 d'encaissement ; brouillon exclu
        first = dict(zip(rows[0], rows[1]))
        self.assertEqual((first["JournalCode"], first["EcritureDate"], first["CompAuxNum"]), ("VT", TODAY.strftime("%Y%m%d"), "C001"))
        self.assertRegex(first["Debit"], r"^\d+,\d{2}$")

    def test_xlsx_export(self):
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(exports.entries_xlsx(LedgerEntry.objects.all()))).active
        self.assertEqual(sheet.max_row, 1 + LedgerEntry.objects.count())

    def test_csv_neutralizes_formulas(self):
        self.assertIn("'=HYPERLINK()", exports.entries_csv(LedgerEntry.objects.all()).decode("utf-8"))


class OpeningEntriesTest(TestCase):
    def test_opening_carries_balance_sheet_and_result(self):
        last = FiscalPeriod.objects.create(name="2025", date_start=date(2025, 1, 1), date_end=date(2025, 12, 31))
        post_entry("VT", date(2025, 6, 1), "Vente", balanced_lines("300.00"), validate=True)
        post_entry("BQ", date(2025, 6, 2), "Frais", [Line(acc("627000"), debit=Decimal("50")),
                                                     Line(acc("512000"), credit=Decimal("50"))], validate=True)
        close_period(last)
        new = FiscalPeriod.objects.create(name="2026", date_start=date(2026, 1, 1), date_end=date(2026, 12, 31))
        entry = generate_opening_entries(last, new)
        lines = {e.account.code: e for e in entry.entries.all()}
        self.assertEqual((lines["411000"].debit, lines["512000"].credit, lines["120000"].credit),
                         (Decimal("300.00"), Decimal("50.00"), Decimal("250.00")))
        self.assertEqual((entry.journal.code, entry.is_validated), ("AN", True))


def fake_document_url(document_type, document_id):
    return f"/pieces/{document_type}/{document_id}/"


def fake_company():
    return {"name": "Hôte SAS", "siren": "987654321"}


def fake_panels(request):
    return [{"label": "Pièces en attente", "value": 3, "url": "/attente/", "alert": True}]


class IntegrationHooksTest(TestCase):
    """Points d'accroche ACCOUNTING : la comptabilité s'intègre à n'importe quel projet."""

    def setUp(self):
        self.client.force_login(accountant())
        self.txn = sale()

    @override_settings(ACCOUNTING={})
    def test_works_standalone(self):
        response = self.client.get(reverse("accounting:dashboard"))
        self.assertTemplateUsed(response, "accounting/standalone_base.html")
        detail = self.client.get(reverse("accounting:entry_detail", args=[self.txn.pk]))
        self.assertContains(detail, "invoice #1")

    @override_settings(ACCOUNTING={"DOCUMENT_URL": "accounting.tests.fake_document_url",
                                   "COMPANY": "accounting.tests.fake_company",
                                   "DASHBOARD_PANELS": ["accounting.tests.fake_panels"]})
    def test_project_hooks(self):
        detail = self.client.get(reverse("accounting:entry_detail", args=[self.txn.pk]))
        self.assertContains(detail, 'href="/pieces/invoice/1/"')
        self.assertContains(self.client.get(reverse("accounting:dashboard")), "Pièces en attente")
        fec = self.client.get(reverse("accounting:exports"), {"download": "fec"})
        self.assertIn("987654321FEC", fec["Content-Disposition"])

    @override_settings(ACCOUNTING={"COMPANY": "accounting.tests.fake_company"})
    def test_settings_override_project_company(self):
        AccountingSettings.objects.filter(pk=1).update(siren="111222333")
        fec = self.client.get(reverse("accounting:exports"), {"download": "fec"})
        self.assertIn("111222333FEC", fec["Content-Disposition"])


@override_settings(ACCOUNTING={})
class AccountingViewsTest(TestCase):
    def setUp(self):
        self.txn = sale()

    def urls(self):
        return [reverse(name, args=args) for name, args in [
            ("accounting:dashboard", []), ("accounting:entry_list", []), ("accounting:entry_create", []),
            ("accounting:entry_detail", [self.txn.pk]), ("accounting:trial_balance", []),
            ("accounting:ledger_index", []), ("accounting:general_ledger", ["411000"]),
            ("accounting:profit_loss", []), ("accounting:vat_report", []), ("accounting:exports", []),
            ("accounting:periods", []),
        ]]

    def test_pages_render_for_accountant(self):
        self.client.force_login(accountant())
        for url in self.urls():
            with self.subTest(url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_users_without_permission_are_refused(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_user(username="x", email="x@example.com", password="x" * 12))
        for url in self.urls():
            with self.subTest(url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_manual_entry_and_draft_edit(self):
        self.client.force_login(accountant())
        base = {"journal": Journal.objects.get(code="OD").pk, "date": TODAY.isoformat(), "reference": "REL-1",
                "description": "Facture fournisseur", "lines-TOTAL_FORMS": "2", "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "1", "lines-MAX_NUM_FORMS": "1000",
                "lines-0-account": acc("622600").pk, "lines-0-debit": "80.00",
                "lines-1-account": acc("401000").pk, "lines-1-credit": "80.00"}
        self.client.post(reverse("accounting:entry_create"), base)
        draft = Transaction.objects.get(reference="REL-1")
        self.assertFalse(draft.is_validated)
        edited = {**base, "description": "Facture corrigée", "lines-0-debit": "90.00", "lines-1-credit": "90.00", "validate": "on"}
        self.client.post(reverse("accounting:entry_edit", args=[draft.pk]), edited)
        txn = Transaction.objects.get(reference="REL-1")
        self.assertEqual((txn.description, txn.amount, txn.is_validated), ("Facture corrigée", Decimal("90.00"), True))

    def test_bank_journal_entry_gets_its_treasury_counterpart(self):
        self.client.force_login(accountant())
        data = {"journal": Journal.objects.get(code="BQ").pk, "date": TODAY.isoformat(), "reference": "REL-2",
                "description": "Frais de tenue de compte", "validate": "on", "lines-TOTAL_FORMS": "1", "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "1", "lines-MAX_NUM_FORMS": "1000",
                "lines-0-account": acc("627000").pk, "lines-0-debit": "12.00"}
        self.client.post(reverse("accounting:entry_create"), data)
        txn = Transaction.objects.get(reference="REL-2")
        self.assertEqual((txn.is_validated, balances(txn)), (True, {"627000": Decimal("12.00"), "512000": Decimal("-12.00")}))

    def test_treasury_account_only_moves_in_its_journal(self):
        with self.assertRaisesMessage(AccountingError, "se mouvemente dans le journal BQ"):
            post_entry("OD", TODAY, "Erreur", [Line(acc("627000"), debit=Decimal("5")), Line(acc("512000"), credit=Decimal("5"))])
        with self.assertRaisesMessage(AccountingError, "se mouvemente dans le journal ST"):
            post_entry("BQ", TODAY, "Erreur", [Line(acc("512000"), debit=Decimal("5")), Line(acc("467100"), credit=Decimal("5"))])
        with self.assertRaisesMessage(AccountingError, "mouvemente son compte de trésorerie"):
            post_entry("BQ", TODAY, "Erreur", [Line(acc("627000"), debit=Decimal("5")), Line(acc("401000"), credit=Decimal("5"))])

    def test_draft_created_outside_posting_is_checked_at_validation(self):
        draft = post_entry("OD", TODAY, "Brouillon", balanced_lines())
        LedgerEntry.objects.filter(transaction=draft, account__code="411000").update(account=acc("512000"))  # admin
        with self.assertRaisesMessage(AccountingError, "se mouvemente dans le journal BQ"):
            validate_entry(draft)
        draft.refresh_from_db()
        self.assertFalse(draft.is_validated)

    def test_mixed_legacy_journal_is_reported(self):
        txn = post_entry("BQ", TODAY, "Encaissement", [Line(acc("512000"), debit=Decimal("5")), Line(acc("411000"), credit=Decimal("5"))])
        LedgerEntry.objects.filter(transaction=txn, account__code="411000").update(account=acc("467100"))  # ancienne écriture
        self.assertEqual(reports.treasury_anomalies(), [{"account": "467100", "owner": "ST", "journal": "BQ", "lines": 1}])

    def test_unbalanced_manual_entry_is_refused(self):
        self.client.force_login(accountant())
        data = {"journal": Journal.objects.get(code="OD").pk, "date": TODAY.isoformat(), "description": "Erreur",
                "lines-TOTAL_FORMS": "2", "lines-INITIAL_FORMS": "0", "lines-MIN_NUM_FORMS": "1", "lines-MAX_NUM_FORMS": "1000",
                "lines-0-account": acc("622600").pk, "lines-0-debit": "80.00",
                "lines-1-account": acc("401000").pk, "lines-1-credit": "70.00"}
        self.assertEqual(self.client.post(reverse("accounting:entry_create"), data).status_code, 200)
        self.assertFalse(Transaction.objects.filter(description="Erreur").exists())

    def test_validate_and_reverse_from_ui(self):
        self.client.force_login(accountant())
        draft = post_entry("OD", TODAY, "Brouillon", balanced_lines())
        self.client.post(reverse("accounting:entry_validate", args=[draft.pk]))
        draft.refresh_from_db()
        self.assertTrue(draft.is_validated)
        self.client.post(reverse("accounting:entry_reverse", args=[draft.pk]))
        self.assertTrue(Transaction.objects.filter(reversal_of=draft).exists())

    def test_downloads(self):
        self.client.force_login(accountant())
        self.assertEqual(self.client.get(reverse("accounting:exports"), {"download": "xlsx"}).status_code, 200)
        self.assertEqual(self.client.get(reverse("accounting:trial_balance"), {"format": "csv"}).status_code, 200)
