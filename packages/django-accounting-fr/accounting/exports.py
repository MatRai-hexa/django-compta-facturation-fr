"""
Exports comptables.

FEC (fichier des écritures comptables, art. L47 A et A47 A-1 du LPF) :
18 colonnes, séparateur tabulation ou barre verticale, dates AAAAMMJJ, montants
avec virgule décimale, écritures validées uniquement, encodage ISO 8859-15.
Nom : <SIREN>FEC<AAAAMMJJ de clôture>.txt. Contrôle conseillé avec l'outil
« Test Compta Demat » de la DGFiP avant remise à l'administration.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Font

from .models import LedgerEntry

FEC_COLUMNS = [
    "JournalCode", "JournalLib", "EcritureNum", "EcritureDate", "CompteNum", "CompteLib", "CompAuxNum",
    "CompAuxLib", "PieceRef", "PieceDate", "EcritureLib", "Debit", "Credit", "EcritureLet", "DateLet",
    "ValidDate", "Montantdevise", "Idevise",
]


# Caractères courants absents de l'ISO 8859-15, remplacés par leur équivalent
_LATIN9 = str.maketrans({"‘": "'", "’": "'", "‚": "'", "“": '"', "”": '"', "„": '"',
                         "–": "-", "—": "-", "−": "-", "…": "...", " ": " ", " ": " ",
                         " ": " ", "•": "-", "·": "-"})


def _latin9(text):
    """Texte encodable en ISO 8859-15 : équivalents usuels, sinon décomposition (ligatures, accents rares), sinon supprimé."""
    out = []
    for char in text.translate(_LATIN9):
        try:
            char.encode("iso-8859-15")
            out.append(char)
        except UnicodeEncodeError:
            out.append(unicodedata.normalize("NFKD", char).encode("iso-8859-15", errors="ignore").decode("iso-8859-15"))
    return "".join(out)


def _fec_text(value, separator):
    text = re.sub(r"[\r\n\t|]+", " ", _latin9(str(value or ""))).replace(separator, " ")
    return " ".join(text.split())


def _fec_date(value):
    return value.strftime("%Y%m%d") if value else ""


def _fec_amount(value):
    return f"{value:.2f}".replace(".", ",")


def fec_lines(start, end):
    return (LedgerEntry.objects.filter(transaction__is_validated=True, transaction__date__gte=start,
                                       transaction__date__lte=end)
            .select_related("transaction", "transaction__journal", "account")
            .order_by("transaction__date", "transaction__validated_at", "transaction_id", "id"))


def fec_file(start, end, siren: str, separator="\t") -> tuple[str, bytes]:
    """Retourne (nom de fichier, contenu) du FEC des écritures validées entre `start` et `end`."""
    out = io.StringIO()
    out.write(separator.join(FEC_COLUMNS) + "\r\n")
    for e in fec_lines(start, end):
        t = e.transaction
        row = [
            t.journal.code, t.journal.label, t.number, _fec_date(t.date), e.account.code, e.account.name,
            e.auxiliary_code, (e.auxiliary_label or e.auxiliary_code), t.reference or t.number, _fec_date(t.piece_date or t.date),
            e.label or t.description, _fec_amount(e.debit), _fec_amount(e.credit), e.reconciliation_ref,
            _fec_date(e.reconciled_at), _fec_date(t.validated_at),
            _fec_amount(e.currency_amount) if e.currency else "", e.currency,
        ]
        out.write(separator.join(_fec_text(v, separator) for v in row) + "\r\n")
    siren = re.sub(r"\D", "", siren or "")[:9] or "000000000"
    return f"{siren}FEC{end:%Y%m%d}.txt", out.getvalue().encode("iso-8859-15")


# === CSV / Excel ===

ENTRY_HEADERS = ["Journal", "Numéro", "Date", "Pièce", "Compte", "Intitulé du compte", "Libellé",
                 "Débit", "Crédit", "Taux TVA", "Validée"]


def _safe(value):
    """Neutralise les formules dans les cellules texte (=, +, -, @ en tête)."""
    text = str(value or "")
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _entry_rows(lines):
    for e in lines.select_related("transaction", "transaction__journal", "account"):
        t = e.transaction
        yield [t.journal.code, t.number or "brouillon", t.date, _safe(t.reference), e.account.code, _safe(e.account.name),
               _safe(e.label), e.debit, e.credit, e.vat_rate if e.vat_rate is not None else "", "oui" if t.is_validated else "non"]


def entries_csv(lines) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(ENTRY_HEADERS)
    for row in _entry_rows(lines):
        writer.writerow([f"{v:%d/%m/%Y}" if hasattr(v, "strftime") else str(v).replace(".", ",") if hasattr(v, "quantize") else v
                         for v in row])
    return ("﻿" + out.getvalue()).encode("utf-8")  # BOM : ouverture correcte dans Excel


def entries_xlsx(lines) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Écritures"
    ws.append(ENTRY_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in _entry_rows(lines):
        ws.append(row)
    for column in ("C",):
        for cell in ws[column][1:]:
            cell.number_format = "DD/MM/YYYY"
    for column in ("H", "I"):
        for cell in ws[column][1:]:
            cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def trial_balance_csv(balance: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(["Compte", "Intitulé", "Débit", "Crédit", "Solde débiteur", "Solde créditeur"])
    fmt = lambda v: f"{v:.2f}".replace(".", ",")
    for group in balance["classes"]:
        for row in group["accounts"]:
            writer.writerow([row["account"].code, _safe(row["account"].name), fmt(row["debit"]), fmt(row["credit"]),
                             fmt(row["balance_debit"]), fmt(row["balance_credit"])])
    writer.writerow(["", "Total", fmt(balance["total"]["debit"]), fmt(balance["total"]["credit"]), "", ""])
    return ("﻿" + out.getvalue()).encode("utf-8")


def analytic_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    fmt = lambda v: f"{v:.2f}".replace(".", ",")
    writer.writerow(["Section", "Libellé", "Produits", "Charges", "Résultat"])
    for row in data["rows"]:
        section = row["section"]
        writer.writerow([section.code if section else "", _safe(section.label) if section else "Non affecté",
                         fmt(row["revenues"]), fmt(row["expenses"]), fmt(row["result"])])
    writer.writerow(["", "Total", fmt(data["revenues"]), fmt(data["expenses"]), fmt(data["result"])])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def annual_accounts_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    fmt = lambda v: f"{v:.2f}".replace(".", ",")
    name, previous = data["period"].name, data["previous"].name if data["previous"] else ""
    writer.writerow(["État", "Rubrique", "Poste", "Brut", "Amortissements et dépréciations", f"Net {name}", f"{previous}"])

    def previous_value(collection, key, field):
        if not collection:
            return ""
        for group in collection:
            for row in group["rows"]:
                if row["key"] == key:
                    return fmt(row[field])
        return ""

    for group in data["balance"]["assets"]:
        for row in group["rows"]:
            writer.writerow(["Actif", group["label"], row["label"], fmt(row["gross"]), fmt(row["contra"]), fmt(row["net"]),
                             previous_value(data["balance_previous"] and data["balance_previous"]["assets"], row["key"], "net")])
    for group in data["balance"]["liabilities"]:
        for row in group["rows"]:
            writer.writerow(["Passif", group["label"], row["label"], "", "", fmt(row["amount"]),
                             previous_value(data["balance_previous"] and data["balance_previous"]["liabilities"], row["key"], "amount")])
    for key, section in data["income"]["sections"].items():
        for line in section["lines"]:
            if line["amount"]:
                writer.writerow(["Compte de résultat", section["label"], line["label"], "", "", fmt(line["amount"]), ""])
    writer.writerow(["Compte de résultat", "", "Résultat net", "", "", fmt(data["income"]["net"]),
                     fmt(data["income_previous"]["net"]) if data["income_previous"] else ""])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def ca3_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    fmt = lambda v: f"{v:.2f}".replace(".", ",")
    writer.writerow(["Ligne", "Libellé", "Base HT", "Taxe"])
    for row in data["operations"]:
        writer.writerow([row["code"], row["label"], fmt(row["base"]), ""])
    for row in data["gross"]:
        writer.writerow([row["code"], row["label"], fmt(row["base"]), fmt(row["tax"])])
    for code, label, value in (("16", "Total de la TVA brute due", data["total_gross"]),
                               ("17", "Dont TVA sur acquisitions intracommunautaires", data["intracom_vat"]),
                               ("19", "TVA déductible sur immobilisations", data["line19"]),
                               ("20", "TVA déductible sur autres biens et services", data["line20"]),
                               ("22", "Report du crédit de la déclaration précédente", data["line22"]),
                               ("23", "Total TVA déductible", data["total_deductible"]),
                               ("25", "Crédit de TVA", data["credit"]), ("28", "TVA nette due", data["due"])):
        writer.writerow([code, label, "", fmt(value)])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def balance_sheet_csv(sheet: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(["Côté", "Rubrique", "Compte", "Intitulé", "Montant"])
    fmt = lambda v: f"{v:.2f}".replace(".", ",")
    for side, sections, total in (("Actif", sheet["assets"], sheet["total_assets"]),
                                  ("Passif", sheet["liabilities"], sheet["total_liabilities"])):
        for section in sections:
            for row in section["accounts"]:
                account = row["account"]
                writer.writerow([side, section["label"], account.code if account else "",
                                 _safe(account.name if account else row["label"]), fmt(row["amount"])])
            writer.writerow([side, f"Total {section['label']}", "", "", fmt(section["total"])])
        writer.writerow([side, f"Total {side.lower()}", "", "", fmt(total)])
    return ("﻿" + out.getvalue()).encode("utf-8")


# === Journaux ===

JOURNAL_HEADERS = ["Journal", "Date", "Numéro", "Pièce", "Compte", "Intitulé du compte", "Compte auxiliaire",
                   "Libellé", "Débit", "Crédit", "Validée"]


def _journal_rows(transactions):
    lines = (LedgerEntry.objects.filter(transaction__in=transactions)
             .select_related("transaction", "transaction__journal", "account")
             .order_by("transaction__date", "transaction__journal__code", "transaction__number", "transaction_id",
                       "-debit", "id"))
    for e in lines.iterator(chunk_size=2000):
        t = e.transaction
        yield [t.journal.code, t.date, t.number or "brouillon", _safe(t.reference), e.account.code, _safe(e.account.name),
               _safe(e.auxiliary_code), _safe(e.label or t.description), e.debit, e.credit, "oui" if t.is_validated else "non"]


def journal_csv(transactions) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(JOURNAL_HEADERS)
    for row in _journal_rows(transactions):
        writer.writerow([f"{v:%d/%m/%Y}" if hasattr(v, "strftime") else str(v).replace(".", ",") if hasattr(v, "quantize")
                         else v for v in row])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def journal_xlsx(transactions, title="Journal") -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31])
    ws.append(JOURNAL_HEADERS)
    for row in _journal_rows(transactions):
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def centralizer_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(["Journal", "Mois", "Écritures", "Débit", "Crédit"])
    amount = lambda v: f"{v:.2f}".replace(".", ",")  # noqa: E731
    for group in data["journals"]:
        for month in group["months"]:
            writer.writerow([group["journal"].code, f"{month['month']:%m/%Y}", month["entries"], amount(month["debit"]),
                             amount(month["credit"])])
        writer.writerow([f"Total {group['journal'].code}", "", group["entries"], amount(group["debit"]), amount(group["credit"])])
    writer.writerow(["Total général", "", data["total"]["entries"], amount(data["total"]["debit"]), amount(data["total"]["credit"])])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def aged_balance_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    writer.writerow(["Compte auxiliaire", "Tiers", *data["labels"], "Total"])
    amount = lambda v: f"{v:.2f}".replace(".", ",")  # noqa: E731
    for row in data["rows"]:
        writer.writerow([_safe(row["code"]), _safe(row["label"]), *map(amount, row["buckets"]), amount(row["total"])])
    writer.writerow(["", "Total", *map(amount, data["totals"]), amount(data["total"])])
    return ("﻿" + out.getvalue()).encode("utf-8")


def bank_state_csv(data: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";")
    amount = lambda v: "" if v is None else f"{v:.2f}".replace(".", ",")  # noqa: E731
    writer.writerow([f"État de rapprochement {data['account'].code} au {data['day']:%d/%m/%Y}"])
    writer.writerow(["Solde comptable", amount(data["book_balance"])])
    writer.writerow(["Écritures non pointées sur le relevé", amount(-data["pending_entries"])])
    writer.writerow(["Opérations du relevé non comptabilisées", amount(data["pending_lines"])])
    writer.writerow(["Solde du relevé attendu", amount(data["expected_bank_balance"])])
    writer.writerow(["Solde du relevé", amount(data["bank_balance"])])
    writer.writerow(["Écart", amount(data["gap"])])
    writer.writerow([])
    writer.writerow(["Nature", "Date", "Pièce", "Libellé", "Montant"])
    for e in data["unmatched_entries"]:
        writer.writerow(["Écriture non pointée", f"{e.transaction.date:%d/%m/%Y}", _safe(e.transaction.reference),
                         _safe(e.label), amount(e.debit - e.credit)])
    for line in data["unmatched_lines"]:
        writer.writerow(["Opération non comptabilisée", f"{line.date:%d/%m/%Y}", _safe(line.reference),
                         _safe(line.label), amount(line.amount)])
    return ("\ufeff" + out.getvalue()).encode("utf-8")
