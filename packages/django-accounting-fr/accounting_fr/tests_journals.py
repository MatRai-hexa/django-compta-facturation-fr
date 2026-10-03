"""Journaux : édition, centralisateur, exports, gestion, journal par moyen de paiement."""
from datetime import date
from decimal import Decimal
import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from . import api, reports
from .models import Account, FiscalPeriod, Journal, PaymentAccount
from .posting import Line, post_entry

CUSTOMER = api.Party("C1", "Camille")


def acc(code):
    return Account.objects.get(code=code)


def user(perms=("view_reports", "export_data", "add_journal", "change_journal", "view_transaction")):
    u = get_user_model().objects.create_user(username="c@x.fr", email="c@x.fr", password="x" * 12)
    u.user_permissions.add(*Permission.objects.filter(content_type__app_label="accounting", codename__in=perms))
    return u



def balances_codes(txn):
    return {e.account.code for e in txn.entries.all()}


@override_settings(ACCOUNTING={})
class JournalsTest(TestCase):
    def setUp(self):
        FiscalPeriod.objects.create(name="2024", date_start=date(2024, 1, 1), date_end=date(2024, 12, 31))
        api.post_sale("s1", date(2024, 1, 10), "F1", "Vente F1", CUSTOMER, [api.SaleLine(Decimal("120.00"))])
        api.post_payment("p1", date(2024, 1, 10), "P1", "Règlement F1", Decimal("120.00"), "stripe", CUSTOMER)
        api.post_sale("s2", date(2024, 2, 3), "F2", "Vente F2", CUSTOMER, [api.SaleLine(Decimal("60.00"))])
        self.draft = post_entry("OD", date(2024, 2, 20), "Brouillon", [Line(acc("627000"), debit=Decimal("5")),
                                                                        Line(acc("401000"), credit=Decimal("5"))])
        self.start, self.end = date(2024, 1, 1), date(2024, 12, 31)

    def test_centralizer_per_journal_and_month(self):
        data = reports.centralizer(self.start, self.end)
        vt = next(g for g in data["journals"] if g["journal"].code == "VT")
        self.assertEqual([(m["month"].month, m["debit"], m["entries"]) for m in vt["months"]],
                         [(1, Decimal("120.00"), 1), (2, Decimal("60.00"), 1)])
        self.assertEqual(data["total"]["debit"], Decimal("305.00"))  # 120 + 60 + 120 + 5
        self.assertTrue(data["is_balanced"] and data["matches_trial_balance"])
        validated = reports.centralizer(self.start, self.end, validated_only=True)
        self.assertEqual(validated["total"]["debit"], Decimal("300.00"))

    def test_journal_book(self):
        vt = Journal.objects.get(code="VT")
        book = reports.journal_book(reports.journal_entries(vt, self.start, self.end))
        self.assertEqual([item["txn"].reference for item in book], ["F1", "F2"])
        first = book[0]
        self.assertEqual([e.account.code for e in first["lines"]], ["411000", "707000", "445711"])  # débits d'abord
        self.assertEqual((first["debit"], first["credit"]), (Decimal("120.00"), Decimal("120.00")))
        totals = reports.journal_totals(vt, self.start, self.end)
        self.assertEqual((totals["entries"], totals["debit"]), (2, Decimal("180.00")))

    def test_pages(self):
        self.client.force_login(user())
        params = {"start": "2024-01-01", "end": "2024-12-31"}
        page = self.client.get(reverse("accounting:journals"), params)
        self.assertContains(page, "Journal centralisateur")
        self.assertContains(page, "totaux identiques à ceux de la balance")
        detail = self.client.get(reverse("accounting:journal_detail", args=["VT"]), params)
        self.assertContains(detail, "Journal VT")
        self.assertContains(detail, "120,00")
        everything = self.client.get(reverse("accounting:journal_all"), params)
        self.assertContains(everything, "Livre-journal")
        self.assertContains(everything, "1 écriture en brouillon")
        printable = self.client.get(reverse("accounting:journal_all"), {**params, "print": "1"})
        self.assertContains(printable, "window.print()")

    def test_exports(self):
        self.client.force_login(user())
        params = {"start": "2024-01-01", "end": "2024-12-31"}
        csv = self.client.get(reverse("accounting:journal_detail", args=["VT"]), {**params, "format": "csv"}).content.decode("utf-8-sig")
        self.assertEqual(csv.splitlines()[0].split(";")[:3], ["Journal", "Date", "Numéro"])
        self.assertIn("VT;10/01/2024;VT2024-00001;F1;411000", csv)
        xlsx = self.client.get(reverse("accounting:journal_all"), {**params, "format": "xlsx"}).content
        self.assertEqual(load_workbook(io.BytesIO(xlsx)).active.max_row, 1 + 3 + 2 + 3 + 2)  # en-tête + lignes
        central = self.client.get(reverse("accounting:journals"), {**params, "format": "csv"}).content.decode("utf-8-sig")
        self.assertIn("Total général;;4;305,00;305,00", central)

    def test_permissions(self):
        self.client.force_login(user(perms=("view_reports",)))
        self.assertEqual(self.client.get(reverse("accounting:journal_all"), {"format": "csv"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("accounting:journal_create")).status_code, 403)
        self.client.force_login(get_user_model().objects.create_user(username="z@x.fr", email="z@x.fr", password="x" * 12))
        self.assertEqual(self.client.get(reverse("accounting:journals")).status_code, 403)

    def test_create_edit_and_delete_journals(self):
        self.client.force_login(user())
        second_bank = Account.objects.create(code="512100", name="Seconde banque", account_type="asset")
        missing = self.client.post(reverse("accounting:journal_create"), {"code": "bq2", "label": "Banque 2", "kind": "bank"})
        self.assertIn("account", missing.context["form"].errors)  # journal de banque sans compte de trésorerie
        taken = self.client.post(reverse("accounting:journal_create"),
                                 {"code": "bq2", "label": "Banque 2", "kind": "bank", "account": acc("512000").pk})
        self.assertIn("account", taken.context["form"].errors)  # déjà le compte de BQ
        response = self.client.post(reverse("accounting:journal_create"),
                                    {"code": "bq2", "label": "Banque 2", "kind": "bank", "account": second_bank.pk})
        self.assertRedirects(response, reverse("accounting:journals"))
        self.assertEqual(Journal.objects.get(code="BQ2").account, second_bank)
        bad = self.client.post(reverse("accounting:journal_create"), {"code": "S T!", "label": "x", "kind": "bank"})
        self.assertIn("code", bad.context["form"].errors)
        # Journal utilisé : code figé, suppression impossible
        self.client.post(reverse("accounting:journal_edit", args=["VT"]), {"code": "XX", "label": "Ventes en ligne", "kind": "sales"})
        self.assertEqual(Journal.objects.get(code="VT").label, "Ventes en ligne")
        self.client.post(reverse("accounting:journal_edit", args=["VT"]), {"delete": "1"})
        self.assertTrue(Journal.objects.filter(code="VT").exists())
        self.client.post(reverse("accounting:journal_edit", args=["BQ2"]), {"delete": "1"})
        self.assertFalse(Journal.objects.filter(code="BQ2").exists())

    def test_payment_method_with_its_own_journal(self):
        payment = api.post_payment("p2", date(2024, 3, 1), "P2", "Règlement F2", Decimal("60.00"), "stripe", CUSTOMER)
        fee = api.post_fee("f2", date(2024, 3, 1), "P2", "Frais", Decimal("1.00"), "stripe")
        self.assertEqual((payment.journal.code, payment.number, fee.number), ("ST", "ST2024-00002", "ST2024-00003"))
        cod = api.post_payment("p3", date(2024, 3, 2), "P3", "Espèces", Decimal("10.00"), "cod", CUSTOMER)
        self.assertEqual((cod.journal.code, balances_codes(cod)), ("CA", {"530000", "411000"}))
        second = Journal.objects.create(code="ST2", label="Stripe 2", kind="bank",
                                        account=Account.objects.create(code="467300", name="Stripe 2", account_type="asset"))
        PaymentAccount.objects.filter(method="stripe").update(journal=second)
        moved = api.post_payment("p4", date(2024, 3, 3), "P4", "Règlement", Decimal("5.00"), "stripe", CUSTOMER)
        self.assertEqual((moved.journal.code, balances_codes(moved)), ("ST2", {"467300", "411000"}))
